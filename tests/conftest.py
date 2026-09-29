from __future__ import annotations

import shutil
from pathlib import Path

import docx
import pytest

from formfill.brains import RuleBrain
from formfill.config import Settings
from formfill.service import FormService

ROOT = Path(__file__).resolve().parent.parent
FORM = ROOT / "forms" / "formular.docx"


@pytest.fixture
def form_path() -> Path:
    return FORM


@pytest.fixture
def form_bytes() -> bytes:
    return FORM.read_bytes()


def fake_converter(docx_path: Path) -> Path:
    """Stand-in for LibreOffice: copies the DOCX so tests stay fast and hermetic."""
    pdf = Path(docx_path).with_suffix(".pdf")
    shutil.copy(docx_path, pdf)
    return pdf


def make_settings(tmp_path: Path) -> Settings:
    return Settings(database_url=f"sqlite:///{tmp_path / 'test.db'}", data_dir=tmp_path,
                    llm_provider="mock", llm_api_key=None, llm_model="test",
                    llm_base_url="http://llm.invalid", llm_timeout=5, max_upload_mb=10)


@pytest.fixture
def settings(tmp_path) -> Settings:
    return make_settings(tmp_path)


@pytest.fixture
def service(settings) -> FormService:
    return FormService(settings, brain=RuleBrain(), converter=fake_converter)


@pytest.fixture
def generic_form(tmp_path) -> Path:
    """A small, unknown form (not the Prihláška) for the generic path."""
    d = docx.Document()
    d.add_heading("Žiadosť o vydanie parkovacej karty", 1)
    d.add_paragraph("Meno a priezvisko:")
    d.add_paragraph("Adresa:")
    d.add_paragraph("EČV vozidla:")
    table = d.add_table(rows=1, cols=2)
    table.cell(0, 0).text = "Dátum:"
    table.cell(0, 1).text = "Podpis:"
    path = tmp_path / "ziadost.docx"
    d.save(path)
    return path
