import docx
import pytest
from docx.oxml.ns import qn

from formfill import document as D
from formfill.devtools.samples import SAMPLES
from formfill.filler import AnchorNotFound, fill_document, find_label
from formfill.fields import Anchor
from formfill.forms import prihlaska


def fill(tmp_path, form_path, name):
    values = dict(SAMPLES[name])
    values.update(prihlaska.derive(values))
    out = tmp_path / f"{name}.docx"
    report = fill_document(form_path, prihlaska.FIELDS, values, out)
    return docx.Document(str(out)), report


def para_with(doc, needle):
    return next(p for p in D.paragraphs(doc) if needle in p.text)


def boxes_after(doc, label):
    para = para_with(doc, label)
    out = []
    for tc in D.cells_after(para.tc):
        if ":" in D.cell_text(tc):  # next label ends the box group
            break
        if D.is_box(tc):
            out.append(D.cell_text(tc).strip())
    return "".join(out)


@pytest.mark.parametrize("name", ["vzor1", "vzor2", "vzor3"])
def test_every_sample_field_is_placed(tmp_path, form_path, name):
    _, report = fill(tmp_path, form_path, name)
    assert report.skipped == []
    assert "email" in report.written


def test_labels_and_values_share_a_line(tmp_path, form_path):
    doc, _ = fill(tmp_path, form_path, "vzor2")
    line = para_with(doc, "Meno:").text
    assert line.index("MARTIN") < line.index("Priezvisko:") < line.index("KOVÁČ")
    assert "Rodné priezvisko:  KOVÁČ" in para_with(doc, "Rodné priezvisko").text
    assert "martin.kovac@example.org" in para_with(doc, "E-mailová adresa").text  # email not uppercased


def test_character_boxes(tmp_path, form_path):
    doc, _ = fill(tmp_path, form_path, "vzor2")
    assert boxes_after(doc, "Dátum narodenia") == "07111996"
    assert boxes_after(doc, "Rodné číslo").replace("/", "") == "9611071234"
    assert boxes_after(doc, "Kód študijného odboru") == "2561M00"


def test_sections_are_disambiguated(tmp_path, form_path):
    doc, _ = fill(tmp_path, form_path, "vzor2")
    streets = [p.text for p in D.paragraphs(doc) if p.text.startswith("Ulica a číslo")]
    assert "HORSKÁ 8" in streets[0]      # trvalé bydlisko
    assert "DUNAJSKÁ 24" in streets[1]   # korešpondenčná adresa
    assert streets[2].strip() == "Ulica a číslo:"  # cudzinec - left empty
    assert "TECHNICKÁ 5" in streets[3]   # stredná škola
    # Protokol on page 6 (for the university) must stay untouched.
    assert all("15.02.2027" not in p.text for p in D.paragraphs(doc) if "komisie" in p.text)


def test_foreigner_scenario(tmp_path, form_path):
    doc, _ = fill(tmp_path, form_path, "vzor3")
    assert boxes_after(doc, "Rodné číslo") == "/"  # no rodné číslo, boxes stay empty
    streets = [p.text for p in D.paragraphs(doc) if p.text.startswith("Ulica a číslo")]
    assert "HLAVNÁ 20" in streets[2]


def test_strike_and_marks(tmp_path, form_path):
    doc, _ = fill(tmp_path, form_path, "vzor2")

    def struck(needle):
        p = para_with(doc, needle).p
        return "".join(D.run_text(r) for r in D.runs(p)
                       if r.find(f"{qn('w:rPr')}/{qn('w:strike')}") is not None)

    assert struck("Súhlasím/Nesúhlasím") == "Súhlasím"  # applicant chose "nesúhlasím"
    assert struck("bakalárske").replace(" ", " ") == "spojené – prvý a druhý stupeň v jednom celku"
    # Gender "muž": X in the box right after the option
    gender = para_with(doc, "Pohlavie")
    assert D.cell_text(D.cells_after(gender.tc)[0]).strip() == "X"
    # School type: X inside the drawn checkbox of "stredná odborná škola"
    cell = next(p.tc for p in D.paragraphs(doc) if p.text.strip() == "stredná odborná škola")
    assert any("X" in D.para_text(p) for p in cell.iter(qn("w:p")))


def test_tables_rows_and_lists(tmp_path, form_path):
    doc, _ = fill(tmp_path, form_path, "vzor2")
    grades = next(t for t in doc.tables if "Ročník" in t._tbl.xml)
    first_row = next(r for r in grades.rows if "MATEMATIKA" in r.cells[0].text)
    assert [c.text for c in first_row.cells][:5] == ["MATEMATIKA", "2", "2", "1", "2"]
    average = para_with(doc, "Priemerný prospech")
    assert D.cell_text(D.cells_after(average.tc)[0]).strip() == "1,75"
    attachments = [p.text.strip() for p in D.paragraphs(doc) if "KÓPIA" in p.text]
    assert attachments == ["KÓPIA MATURITNÉHO VYSVEDČENIA", "KÓPIA VŠ DIPLOMU"]
    assert "2027" in para_with(doc, "Akademický rok").text and "2028" in para_with(doc, "Akademický rok").text


def test_current_maturant_fills_school_code(tmp_path, form_path):
    doc, _ = fill(tmp_path, form_path, "vzor1")
    assert boxes_after(doc, "Kód strednej školy") == "790200000"


def test_template_is_not_modified(tmp_path, form_path):
    before = form_path.read_bytes()
    fill(tmp_path, form_path, "vzor1")
    assert form_path.read_bytes() == before


def test_find_label_rules(form_path):
    paras = D.paragraphs(docx.Document(str(form_path)))
    para, m = find_label(paras, Anchor("Priezvisko"))
    assert para.text[m.start():m.end()] == "Priezvisko:"  # not "Rodné priezvisko"
    with pytest.raises(AnchorNotFound):
        find_label(paras, Anchor("Neexistujúce pole"))
