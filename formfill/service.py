"""Application service: glues upload -> form detection -> conversation -> PDF rendering,
and persists everything so a conversation survives an application restart."""

from __future__ import annotations

import io
import logging
import threading
import uuid
import zipfile
from collections import defaultdict
from pathlib import Path
from typing import Callable

import docx

from .brains import LLMBrain, RuleBrain
from .config import Settings
from .db import FormSession, Message, make_sessionmaker
from .engine import DONE, ConversationEngine, ConversationState
from .fields import schema_from_json, schema_to_json
from .filler import fill_document
from .forms import registry
from .llm.client import OpenAICompatibleChat
from .pdf import ConversionError, docx_to_pdf

log = logging.getLogger(__name__)


class NotFound(LookupError):
    pass


class InvalidDocument(ValueError):
    pass


def make_brain(settings: Settings):
    if settings.llm_provider == "mock":
        return RuleBrain()
    chat = OpenAICompatibleChat(settings.llm_api_key, settings.llm_model, settings.llm_base_url,
                                timeout=settings.llm_timeout)
    return LLMBrain(chat)


class FormService:
    def __init__(self, settings: Settings, brain=None,
                 converter: Callable[[Path], Path] = docx_to_pdf):
        self.settings = settings
        self.brain = brain or make_brain(settings)
        self.converter = converter
        settings.data_dir.mkdir(parents=True, exist_ok=True)
        self.Session = make_sessionmaker(settings.database_url)
        self._locks: dict[str, threading.Lock] = defaultdict(threading.Lock)

    # ---------------------------------------------------------------- helpers
    def _engine(self, s: FormSession) -> ConversationEngine:
        return ConversationEngine(s.title, s.description, schema_from_json(s.schema),
                                  registry.derive_for(s.form_id or ""), self.brain)

    def _load(self, db, session_id: str) -> FormSession:
        s = db.get(FormSession, session_id)
        if s is None:
            raise NotFound(session_id)
        return s

    def describe(self, s: FormSession) -> dict:
        engine = self._engine(s)
        state = ConversationState.from_dict(s.state)
        return {
            "id": s.id, "title": s.title, "filename": s.filename, "form_id": s.form_id,
            "phase": s.phase, "created_at": s.created_at.isoformat(),
            "updated_at": s.updated_at.isoformat(), "progress": engine.progress(state),
            "pdf_ready": bool(s.pdf_path and Path(s.pdf_path).exists()),
        }

    # ------------------------------------------------------------------ API
    def create_session(self, filename: str, content: bytes) -> dict:
        if len(content) > self.settings.max_upload_mb * 1024 * 1024:
            raise InvalidDocument("Súbor je príliš veľký.")
        try:
            zipfile.ZipFile(io.BytesIO(content))
            docx.Document(io.BytesIO(content))
        except Exception as exc:  # noqa: BLE001 - any parse failure means "not a docx"
            raise InvalidDocument("Súbor nie je platný dokument .docx.") from exc

        session_id = str(uuid.uuid4())
        template = self.settings.data_dir / "uploads" / f"{session_id}.docx"
        template.parent.mkdir(parents=True, exist_ok=True)
        template.write_bytes(content)

        known = registry.detect_known_form(registry.document_text(template))
        if known:
            form_id, title, description, fields = known.form_id, known.title, known.description, known.fields
        else:
            form_id = None
            title, description, fields = self.brain.derive_fields(registry.outline(template))
            if not fields:
                template.unlink(missing_ok=True)
                raise InvalidDocument("V dokumente som nenašiel žiadne polia na vyplnenie "
                                      "(napr. „Meno: ____“). Nahrajte prosím tlačivo.")

        s = FormSession(id=session_id, filename=filename, template_path=str(template),
                        form_id=form_id, title=title, description=description,
                        schema=schema_to_json(fields), state={})
        engine = self._engine(s)
        state = ConversationState()
        reply = engine.start(state)
        s.state, s.phase = state.to_dict(), state.phase
        s.messages.append(Message(role="assistant", content=reply.text, source=reply.source))
        with self.Session() as db, db.begin():
            db.add(s)
        return {"session": self.describe(s), "reply": reply.text}

    def list_sessions(self) -> list[dict]:
        with self.Session() as db:
            rows = db.query(FormSession).order_by(FormSession.updated_at.desc()).all()
            return [self.describe(s) for s in rows]

    def get_session(self, session_id: str) -> dict:
        with self.Session() as db:
            s = self._load(db, session_id)
            state = ConversationState.from_dict(s.state)
            return {**self.describe(s),
                    "values": {k: v for k, v in state.values.items() if not k.startswith("_")},
                    "messages": [{"role": m.role, "content": m.content,
                                  "created_at": m.created_at.isoformat()} for m in s.messages]}

    def send_message(self, session_id: str, text: str) -> dict:
        text = text.strip()
        if not text:
            raise ValueError("Správa je prázdna.")
        with self._locks[session_id], self.Session() as db, db.begin():
            s = self._load(db, session_id)
            engine = self._engine(s)
            state = ConversationState.from_dict(s.state)
            history = [{"role": m.role, "content": m.content} for m in s.messages]
            s.messages.append(Message(role="user", content=text))
            reply = engine.handle(state, history, text)
            reply_text = reply.text
            if reply.ready:
                try:
                    skipped = self._render(s, engine, state)
                    reply_text += ("\n\n✅ Hotovo! Vyplnené tlačivo stiahnete tlačidlom „Stiahnuť PDF“ "
                                   "vpravo hore (v CLI sa stiahne automaticky).")
                    if skipped:
                        reply_text += "\n\nNiektoré údaje sa nepodarilo umiestniť: " + "; ".join(
                            f"{k} ({why})" for k, why in skipped)
                except ConversionError as exc:
                    log.error("PDF rendering failed: %s", exc)
                    state.phase = "review"
                    reply_text += f"\n\n❌ Nepodarilo sa vygenerovať PDF ({exc}). Skúste to znova napísaním „áno“."
            s.state, s.phase = state.to_dict(), state.phase
            s.messages.append(Message(role="assistant", content=reply_text, source=reply.source))
            return {"reply": reply_text, "session": self.describe(s), "errors": reply.errors}

    def _render(self, s: FormSession, engine: ConversationEngine, state: ConversationState):
        out_dir = self.settings.data_dir / "outputs" / s.id
        filled = out_dir / "vyplnene_tlacivo.docx"
        report = fill_document(s.template_path, engine.fields, state.values, filled)
        s.pdf_path = str(self.converter(filled))
        return report.skipped

    def render_pdf(self, session_id: str) -> Path:
        """Re-render on demand (e.g. template or filler improved since)."""
        with self._locks[session_id], self.Session() as db, db.begin():
            s = self._load(db, session_id)
            if s.phase != DONE:
                raise ValueError("Formulár ešte nie je potvrdený.")
            self._render(s, self._engine(s), ConversationState.from_dict(s.state))
            return Path(s.pdf_path)

    def pdf_path(self, session_id: str) -> Path:
        with self.Session() as db:
            s = self._load(db, session_id)
            if not s.pdf_path or not Path(s.pdf_path).exists():
                raise NotFound("PDF zatiaľ nie je vygenerované.")
            return Path(s.pdf_path)

    def delete_session(self, session_id: str) -> None:
        with self.Session() as db, db.begin():
            s = self._load(db, session_id)
            for path in (s.template_path, s.pdf_path):
                if path:
                    Path(path).unlink(missing_ok=True)
            db.delete(s)
