"""Form detection: known form definitions first, LLM-derived schema for anything else."""

from __future__ import annotations

import re
from dataclasses import dataclass, field
from pathlib import Path
from typing import Callable

import docx

from .. import document as D
from ..brains import DocumentOutline
from ..fields import Field
from . import prihlaska

LABEL_RE = re.compile(r"(?<![\w.])([A-ZÁ-Ž][^:\n\t]{0,60}?)\s*(?:\d{1,2}\s*\))?\s*:")


@dataclass
class FormDefinition:
    form_id: str
    title: str
    description: str
    fields: list[Field]
    derive: Callable[[dict], dict] = field(default=lambda values: {})


KNOWN_FORMS = {
    prihlaska.FORM_ID: FormDefinition(prihlaska.FORM_ID, prihlaska.TITLE, prihlaska.DESCRIPTION,
                                      prihlaska.FIELDS, prihlaska.derive),
}
DETECTORS = {prihlaska.FORM_ID: prihlaska.DETECT_PHRASES}


def document_text(path: str | Path) -> str:
    doc = docx.Document(str(path))
    return "\n".join(p.text for p in D.paragraphs(doc))


def detect_known_form(text: str) -> FormDefinition | None:
    flat = text.replace(D.NBSP, " ")
    for form_id, phrases in DETECTORS.items():
        if all(phrase in flat for phrase in phrases):
            return KNOWN_FORMS[form_id]
    return None


def derive_for(form_id: str) -> Callable[[dict], dict]:
    form = KNOWN_FORMS.get(form_id)
    return form.derive if form else (lambda values: {})


def outline(path: str | Path, max_slots: int = 120) -> DocumentOutline:
    """Label slots of an unknown form: "Label:" occurrences, with occurrence index and the
    nearest heading-like context so duplicated labels ("Ulica:") can be told apart."""
    doc = docx.Document(str(path))
    paras = D.paragraphs(doc)
    slots, counts, context = [], {}, ""
    for para in paras:
        text = para.text
        labels = list(LABEL_RE.finditer(text))
        if not labels and text.strip():
            context = text.strip()[:80]
        for m in labels:
            label = re.sub(r"\s+", " ", m.group(1)).strip()
            nth = counts.get(label, 0)
            counts[label] = nth + 1
            slots.append({"id": len(slots), "label": label, "nth": nth, "context": context})
            if len(slots) >= max_slots:
                break
    return DocumentOutline(text="\n".join(p.text for p in paras)[:6000], slots=slots)
